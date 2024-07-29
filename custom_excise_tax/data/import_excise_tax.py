import openpyxl
import requests
from io import BytesIO
from odoo import api, SUPERUSER_ID

def import_excise_tax(cr, registry):
    env = api.Environment(cr, SUPERUSER_ID, {})
    ExciseTax = env['excise.tax']

    # URL of the Excel file
    url = "https://stallesbon.blob.core.windows.net/public/excise_tax_data.xlsx"

    # Download the file
    response = requests.get(url)
    response.raise_for_status()  # Ensure we notice bad responses

    # Load the workbook from the downloaded file
    workbook = openpyxl.load_workbook(filename=BytesIO(response.content))
    #workbook = openpyxl.load_workbook('/odoo/custom/brewsoft/custom_excise_tax/data/excise_tax_data.xlsx')
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        name, excise_tax, special_excise_tax, packaging_tax = row
        ExciseTax.create({
            'name': name,
            'excise_tax': excise_tax,
            'special_excise_tax': special_excise_tax,
            'packaging_tax': packaging_tax,
        })
